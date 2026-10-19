"""Unit and Integration Tests for Verism AI Backend Service."""

import pytest
from app.core.security import ferpa_cipher, verify_password, get_password_hash, create_access_token, decode_access_token
from app.services.ferpa_service import ferpa_service
from app.services.gemini_service import gemini_service
from app.services.vector_service import vector_service
from app.schemas.grading import StructuredGradingOutput


def test_ferpa_encryption_and_decryption():
    """Verify that student PII is safely encrypted and decrypted accurately."""
    original_name = "Rahul Sharma"
    encrypted = ferpa_cipher.encrypt(original_name)

    assert encrypted != original_name
    assert len(encrypted) > 10

    decrypted = ferpa_cipher.decrypt(encrypted)
    assert decrypted == original_name


def test_ferpa_pseudonymization_and_sanitization():
    """Verify that student PII is masked and pseudonymized for AI prompts."""
    sub_id = "sub_123"
    name = "Rahul Sharma"
    roll = "2024-047"
    email = "rahul.sharma@university.edu"
    raw_text = "Submitted by Rahul Sharma (2024-047, rahul.sharma@university.edu). Here is my graph algorithm implementation."

    pseudo_id, sanitized = ferpa_service.prepare_for_ai_evaluation(
        submission_id=sub_id,
        student_name=name,
        student_roll=roll,
        student_email=email,
        text_content=raw_text,
    )

    assert pseudo_id.startswith("STUDENT_")
    assert name not in sanitized
    assert roll not in sanitized
    assert email not in sanitized
    assert "[STUDENT_NAME]" in sanitized
    assert "[STUDENT_ROLL]" in sanitized
    assert "[STUDENT_EMAIL]" in sanitized


def test_password_hashing():
    """Verify password hashing and timing-safe verification."""
    password = "secret_password_123"
    hashed = get_password_hash(password)

    assert hashed != password
    assert verify_password(password, hashed) is True
    assert verify_password("wrong_password", hashed) is False


def test_jwt_token_handling():
    """Verify JWT token creation and decoding."""
    token = create_access_token("inst_priya_rao")
    assert token is not None

    decoded = decode_access_token(token)
    assert decoded is not None
    assert decoded.get("sub") == "inst_priya_rao"


def test_gemini_structured_fallback_evaluation():
    """Verify that rubric evaluation produces schema-compliant StructuredGradingOutput."""
    sample_text = "This report applies inheritance and polymorphism to model a banking transaction ledger."
    rubric_criteria = [
        {"name": "Understanding", "max_score": 30.0, "weight": 30.0},
        {"name": "Accuracy", "max_score": 30.0, "weight": 30.0},
        {"name": "Methodology", "max_score": 20.0, "weight": 20.0},
        {"name": "Presentation", "max_score": 20.0, "weight": 20.0},
    ]

    result = gemini_service.evaluate_submission(
        submission_text=sample_text,
        rubric_criteria=rubric_criteria,
        assignment_title="Assignment 3",
    )

    assert isinstance(result, StructuredGradingOutput)
    assert 0.0 <= result.overall_score <= 100.0
    assert len(result.criteria_evaluations) == 4
    assert result.feedback_summary is not None
    assert isinstance(result.ai_generated_probability, float)


def test_vector_similarity_and_matching_phrases():
    """Verify vector embedding generation and matching phrase detection."""
    text_a = "This report applies inheritance and polymorphism to model a banking transaction ledger, tracing how an abstract Account base class branches into Savings and Current subtypes."
    text_b = "The solution uses an abstract Account base class branches into Savings and Current subtypes, then overrides the interest method for each account type."

    matches = vector_service.find_matching_phrases(text_a, text_b, min_words=4)
    assert len(matches) > 0

    vec_a = vector_service.generate_embedding(text_a)
    assert len(vec_a) == vector_service.EMBEDDING_DIM


def test_supabase_client_and_service():
    """Verify Supabase client and SupabaseService PostgREST query layer."""
    from app.core.supabase import supabase
    from app.services.supabase_service import supabase_service

    assert supabase is not None
    assert supabase_service is not None

    courses = supabase_service.get_courses()
    assert isinstance(courses, list)
    assert len(courses) >= 4


def test_supabase_ferpa_submission_lifecycle():
    """Verify SupabaseService FERPA encryption and audit logging."""
    from app.services.supabase_service import supabase_service

    # Create submission with encrypted PII
    sub = supabase_service.create_submission(
        assignment_id="assign_oop_a3",
        student_name="Test Student",
        student_roll="2024-999",
        file_name="test_student_project.pdf",
        extracted_text="Testing tree and graph search algorithms.",
        student_email="test.student@university.edu",
    )
    assert sub is not None
    assert sub["id"].startswith("sub_")

    # Audit log entry
    audit = supabase_service.log_audit(
        actor_name="Dr. Priya Rao",
        action="TEST_ACTION",
        entity_type="Submission",
        entity_id=sub["id"],
        details={"test": True},
    )
    assert audit is not None
    assert audit["action"] == "TEST_ACTION"


def test_fastapi_health_endpoint():
    """Verify FastAPI /health endpoint returns accurate diagnostic status."""
    from fastapi.testclient import TestClient
    from app.main import app

    client = TestClient(app)
    response = client.get("/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "healthy"
    assert data["supabase_configured"] is True
    assert data["ferpa_encryption"] == "active"
    assert data["qdrant_configured"] is True
    assert data["qdrant_mode"] in ("embedded/in-memory", "http://localhost:6333", ":memory:")


def test_user_signup_and_login_api():
    """Verify user signup and login endpoints work with Supabase/PostgreSQL sync."""
    import uuid
    from fastapi.testclient import TestClient
    from app.main import app

    client = TestClient(app)
    unique_email = f"prof_{uuid.uuid4().hex[:6]}@university.edu"

    # Test Sign Up
    signup_res = client.post(
        "/api/v1/auth/signup",
        json={
            "name": "Prof. Alan Turing",
            "email": unique_email,
            "password": "SecurePassword123!",
            "role": "Professor · Computer Science",
        },
    )
    assert signup_res.status_code == 201
    signup_data = signup_res.json()
    assert "access_token" in signup_data
    assert signup_data["user"]["email"] == unique_email
    token = signup_data["access_token"]

    # Test Get Current User /me with Bearer token
    me_res = client.get(
        "/api/v1/auth/me",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert me_res.status_code == 200
    assert me_res.json()["name"] == "Prof. Alan Turing"

    # Test Login
    login_res = client.post(
        "/api/v1/auth/login",
        json={
            "email": unique_email,
            "password": "SecurePassword123!",
        },
    )
    assert login_res.status_code == 200
    assert "access_token" in login_res.json()


def test_evaluation_with_reference_assignment_and_web_plagiarism():
    """Verify Gemini evaluator incorporates reference assignments and reports 3-way plagiarism metrics."""
    rubric = [
        {"name": "Correctness", "max_score": 50.0, "weight": 50.0, "description": "Correct algorithmic solution"},
        {"name": "Code Style", "max_score": 50.0, "weight": 50.0, "description": "Clean modular implementation"},
    ]
    reference = "def dijkstra(graph, start): heap = [(0, start)] ... optimal benchmark code"
    student_submission = "def dijkstra(g, s): q = [(0, s)] ... student solution"

    result = gemini_service.evaluate_submission(
        submission_text=student_submission,
        rubric_criteria=rubric,
        assignment_title="Dijkstra Shortest Path",
        reference_assignment_text=reference,
    )

    assert result.overall_score >= 0.0
    assert len(result.criteria_evaluations) == 2
    assert result.ai_generated_probability >= 0.0
    assert result.peer_to_peer_plagiarism_pct is not None
    assert result.web_plagiarism_pct is not None
    assert isinstance(result.web_sources, list)


def test_file_upload_database_persistence_and_retrieval():
    """Verify uploaded files are saved in the database with file_content and extracted_text."""
    from fastapi.testclient import TestClient
    from app.main import app
    from app.core.database import SessionLocal
    from app.models.submission import Submission
    from app.models.assignment import Assignment
    from app.models.course import Course

    client = TestClient(app)
    db = SessionLocal()

    # Ensure a test course and assignment exist
    try:
        course = db.query(Course).first()
        if not course:
            course = Course(id="c_test_upload", name="Test Course", code="TC101", section="A")
            db.add(course)
            db.commit()

        assignment = db.query(Assignment).filter(Assignment.course_id == course.id).first()
        if not assignment:
            assignment = Assignment(id="assign_test_upload", course_id=course.id, title="Test Assignment")
            db.add(assignment)
            db.commit()
            db.refresh(assignment)

        # Upload a code file
        code_content = b"def bubble_sort(arr):\n    for i in range(len(arr)):\n        for j in range(len(arr)-i-1):\n            if arr[j] > arr[j+1]: arr[j], arr[j+1] = arr[j+1], arr[j]\n    return arr\n"
        upload_res = client.post(
            f"/api/v1/assignments/{assignment.id}/submissions/batch-upload",
            files=[("files", ("sharma_rahul_sort.py", code_content, "text/x-python"))],
        )
        assert upload_res.status_code == 200, upload_res.text
        upload_data = upload_res.json()
        assert upload_data["total_files"] == 1
        sub_id = upload_data["submissions"][0]["id"]

        # Verify database record
        saved_sub = db.query(Submission).filter(Submission.id == sub_id).first()
        assert saved_sub is not None
        assert saved_sub.student_name == "Rahul Sharma"
        assert saved_sub.file_name == "sharma_rahul_sort.py"
        assert saved_sub.extracted_text is not None
        assert "bubble_sort" in saved_sub.extracted_text
        assert saved_sub.file_content is not None
        assert "bubble_sort" in saved_sub.file_content

        # Verify API list returns decrypted student and file info
        list_res = client.get(f"/api/v1/assignments/{assignment.id}/submissions")
        assert list_res.status_code == 200
        items = list_res.json()
        matching = [item for item in items if item["id"] == sub_id]
        assert len(matching) == 1
        assert matching[0]["student"] == "Rahul Sharma"
        assert matching[0]["file_name"] == "sharma_rahul_sort.py"
    finally:
        db.close()


def test_docx_text_extraction_paragraphs_and_tables():
    """Verify python-docx correctly extracts both paragraphs and table cells."""
    import io
    import docx
    from app.services.pdf_service import pdf_service

    doc = docx.Document()
    doc.add_heading("Research Report: Distributed Consensus", level=1)
    doc.add_paragraph("This paper analyzes Paxos and Raft consensus protocols under high network latency.")
    doc.add_paragraph("Section 2: Comparative Throughput Analysis.")

    table = doc.add_table(rows=2, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Protocol"
    hdr_cells[1].text = "Leader Election (ms)"
    hdr_cells[2].text = "Message Complexity"

    row_cells = table.rows[1].cells
    row_cells[0].text = "Raft"
    row_cells[1].text = "150ms"
    row_cells[2].text = "O(n)"

    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    extracted_text = pdf_service.extract_text_from_bytes(docx_bytes, "distributed_consensus.docx")
    assert extracted_text is not None
    assert "Distributed Consensus" in extracted_text
    assert "Paxos and Raft consensus" in extracted_text
    assert "Comparative Throughput Analysis" in extracted_text
    assert "Protocol | Leader Election (ms) | Message Complexity" in extracted_text
    assert "Raft | 150ms | O(n)" in extracted_text


def test_docx_submission_upload_endpoint():
    """Verify multipart .docx submission uploads parse text and save to DB."""
    import io
    import docx
    from fastapi.testclient import TestClient
    from app.main import app
    from app.core.database import SessionLocal
    from app.models.assignment import Assignment
    from app.models.course import Course
    from app.models.submission import Submission

    client = TestClient(app)
    db = SessionLocal()
    try:
        course = db.query(Course).first()
        if not course:
            course = Course(id="c_docx_test", name="Docx Course", code="DC101")
            db.add(course)
            db.commit()

        assignment = db.query(Assignment).filter(Assignment.course_id == course.id).first()
        if not assignment:
            assignment = Assignment(id="assign_docx_test", course_id=course.id, title="Docx Assignment")
            db.add(assignment)
            db.commit()

        # Build in-memory docx
        doc = docx.Document()
        doc.add_paragraph("Student analysis on deep learning transformers and attention mechanisms.")
        bio = io.BytesIO()
        doc.save(bio)
        docx_bytes = bio.getvalue()

        # Upload via /api/v1/assignments/{assignment_id}/submissions
        res = client.post(
            f"/api/v1/assignments/{assignment.id}/submissions",
            files={"file": ("deep_learning_essay.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"student_name": "Maya Patel", "student_roll": "2024-555"},
        )
        assert res.status_code == 200, res.text
        data = res.json()
        assert data["student"] == "Maya Patel"
        assert data["file_name"] == "deep_learning_essay.docx"
        assert data["status"] == "pending"

        sub_id = data["id"]
        saved = db.query(Submission).filter(Submission.id == sub_id).first()
        assert saved is not None
        assert saved.student_name == "Maya Patel"
        assert saved.extracted_text is not None
        assert "deep learning transformers" in saved.extracted_text
        assert saved.content is not None
        assert "deep learning transformers" in saved.content
    finally:
        db.close()


def test_gemini_schema_bidirectional_mapping():
    """Verify StructuredGradingOutput models validate all required audit keys."""
    from app.schemas.grading import (
        StructuredGradingOutput,
        GradingSummary,
        RubricBreakdownItem,
        IntegrityAudit,
        PeerPlagiarismAudit,
        WebPlagiarismAudit,
        WebSourceItem,
        AIContentAssessment,
        FeedbackDetails,
    )

    data = {
        "grading_summary": {
            "score": 88.0,
            "max_points": 100.0,
            "percentage": 88.0,
            "grade": "B+",
        },
        "rubric_breakdown": [
            {
                "criterion_name": "Technical Depth",
                "awarded_points": 28.0,
                "max_points": 30.0,
                "rationale": "Clear algorithmic analysis with rigorous asymptotic bounds.",
            }
        ],
        "integrity_audit": {
            "peer_plagiarism": {
                "similarity_percentage": 5.2,
                "matched_peer_id": "STUDENT_A",
                "matching_excerpts": ["standard library initialization"],
            },
            "web_plagiarism": {
                "similarity_percentage": 12.0,
                "matched_sources": [
                    {
                        "url": "https://en.wikipedia.org/wiki/Quicksort",
                        "title": "Quicksort - Wikipedia",
                        "snippet": "Quicksort is an efficient divide-and-conquer sorting algorithm.",
                        "match_pct": 12.0,
                    }
                ],
            },
            "ai_content_assessment": {
                "likelihood_rating": "low",
                "ai_probability": 8.5,
                "stylistic_indicators": ["Varied sentence structure and idiosyncratic code formatting."],
            },
        },
        "feedback": {
            "strengths": ["Strong theoretical foundation", "Clear diagrams"],
            "actionable_corrections": ["Add empirical benchmarks for non-uniform data."],
            "summary": "Solid submission demonstrating good domain grasp.",
        },
    }

    output = StructuredGradingOutput(**data)
    assert output.overall_score == 88.0
    assert output.grading_summary.percentage == 88.0
    assert output.integrity_audit.peer_plagiarism.similarity_percentage == 5.2
    assert output.integrity_audit.web_plagiarism.similarity_percentage == 12.0
    assert len(output.integrity_audit.web_plagiarism.matched_sources) == 1
    assert output.integrity_audit.web_plagiarism.matched_sources[0].url == "https://en.wikipedia.org/wiki/Quicksort"
    assert output.integrity_audit.ai_content_assessment.ai_probability == 8.5
    assert len(output.feedback.strengths) == 2


def test_assignment_creation_and_supabase_sync():
    """Verify /assignments endpoint creates assignment with UUID, syncs course, and sets up evaluation."""
    from fastapi.testclient import TestClient
    from app.main import app
    from app.core.database import SessionLocal
    from app.models.assignment import Assignment

    client = TestClient(app)
    res = client.post(
        "/api/v1/assignments",
        json={
            "course_id": "c_cs201",
            "title": "CS201 · Data Structures - Section A Assignment",
            "description": "Sorting algorithms and tree traversal implementation.",
            "reference_assignment_text": "Reference Exemplar: Dijkstra and Quicksort with complexity analysis.",
        },
    )
    assert res.status_code == 200, res.text
    data = res.json()
    assign_id = data["id"]
    assert assign_id.startswith("assign_")
    assert data["title"] == "CS201 · Data Structures - Section A Assignment"

    # Verify local DB persistence
    db = SessionLocal()
    try:
        assign = db.query(Assignment).filter(Assignment.id == assign_id).first()
        assert assign is not None
        assert assign.reference_assignment_text is not None

        # Trigger evaluation on this assignment ID
        eval_res = client.post(
            f"/api/v1/assignments/{assign_id}/evaluate",
            json={"similarity_threshold": 90.0},
        )
        assert eval_res.status_code == 200, eval_res.text
        eval_data = eval_res.json()
        assert eval_data["status"] in ("queued", "processing", "completed")
    finally:
        db.close()


def test_evaluation_fallback_for_unregistered_assignment_id():
    """Verify that evaluating an unregistered/temporary assignment ID dynamically creates it without crashing."""
    from fastapi.testclient import TestClient
    from app.main import app
    import uuid

    client = TestClient(app)
    temp_id = f"a_{uuid.uuid4().hex[:8]}"

    # Evaluating a non-existent assignment ID should self-heal and not throw 404/500
    eval_res = client.post(
        f"/api/v1/assignments/{temp_id}/evaluate",
        json={
            "title": "Dynamic Fallback Assignment",
            "reference_assignment_text": "Dynamic benchmark solution.",
            "similarity_threshold": 90.0,
        },
    )
    assert eval_res.status_code == 200, eval_res.text
    eval_data = eval_res.json()
    assert eval_data["status"] in ("queued", "processing", "completed")




